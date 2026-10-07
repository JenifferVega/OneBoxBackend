"""Builds the starting state and the fake inputs of the conversation tests.

    python tests/planning/conversations/_build_fixtures.py

Starting states come from the plans the reanalysis cases produced
(cases/*/runs/latest/after.json), with readable ids instead of 'new-N' and a
few curated changes, so a script can say "the task bion-posts" and mean it.

Fake inputs follow the REAL shapes:
  * Trello: exactly what the Trello REST API returns for the two reads OneBox
    makes (api/services/trello.py):
        GET /1/boards/{board}/lists?fields=name,closed
        GET /1/lists/{list}/cards?fields=name,desc,due,idList,dateLastActivity,closed
    Cards exported by OneBox carry the description agent/tools/trello.py
    writes ("Owner (OneBox): X" ... "OneBox task: <taskId>"), and the tasks
    carry trelloCardId / trelloCardListId / trelloSyncHash like after a push.
  * Documents: a .docx (read by agent.document_parser.extract_text, the same
    extractor production uses) and a .md.
Re-running overwrites the files; the scripts (script.json) are hand-written.
"""
import copy
import hashlib
import json
from pathlib import Path

HERE = Path(__file__).resolve().parent
CASES = HERE.parent / "cases"


def _load(case):
    return json.loads((CASES / case / "runs" / "latest" / "after.json").read_text(encoding="utf-8"))


def _rename(export, mapping, drop=()):
    ex = copy.deepcopy(export)
    ex["tasks"] = [t for t in ex["tasks"] if t["taskId"] not in drop]
    for t in ex["tasks"]:
        t["taskId"] = mapping.get(t["taskId"], t["taskId"])
        t["dependsOn"] = [mapping.get(d, d) for d in t.get("dependsOn") or [] if d not in drop]
        t.setdefault("createdBy", "IA")
    for i in ex.get("insights") or []:
        if i.get("taskId"):
            i["taskId"] = mapping.get(i["taskId"], i["taskId"])
    ex.pop("events", None)
    ex.pop("applyErrors", None)
    ex["attachments"] = []          # conversations start from the plan, not from the sources
    return ex


def _task(ex, tid):
    return next(t for t in ex["tasks"] if t["taskId"] == tid)


def _write(folder, name, data):
    p = HERE / folder / name
    p.parent.mkdir(parents=True, exist_ok=True)
    if isinstance(data, (dict, list)):
        p.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    elif isinstance(data, bytes):
        p.write_bytes(data)
    else:
        p.write_text(data, encoding="utf-8")


# ── BION ─────────────────────────────────────────────────────────────────────
BION_IDS = {"new-1": "bion-dropbox", "new-2": "bion-banner", "new-3": "bion-posts", "new-4": "bion-urls",
            "new-5": "bion-videos"}


def bion_base():
    ex = _rename(_load("bion_reanalyze"), BION_IDS)
    ex["project"]["participants"] = [
        {"name": "Jesus", "role": "PM", "email": "ventas.javc@gmail.com", "phone": ""},
        {"name": "Claudia", "role": "Ads Lead", "email": "claudiamchimento@gmail.com", "phone": ""},
        {"name": "Javier Costa", "role": "Cuentas / comunicación cliente", "email": "javier.costa@twozero.example", "phone": ""},
        {"name": "Rafael Sardinha", "role": "Cliente BION (accesos)", "email": "rafael@sportsbion.example", "phone": "+351910000001"},
        {"name": "Raquel", "role": "Cliente BION (redes)", "email": "", "phone": "+351910000002"},
    ]
    ex["project"]["planVersion"] = 3
    _task(ex, "bion-posts")["assignedTo"] = "Raquel"
    _task(ex, "bion-urls")["assignedTo"] = "Rafael Sardinha"
    return ex


def _card_desc(task):
    parts = []
    if task.get("assignedTo"):
        parts.append(f"Owner (OneBox): {task['assignedTo']}")
    if task.get("status") == "blocked":
        parts.append("Status in OneBox: BLOCKED")
    parts.append(f"OneBox task: {task['taskId']}")
    return "\n\n".join(parts)


def _hash(name, desc, due):
    return hashlib.sha256("\x00".join((name or "", desc or "", due or "")).encode("utf-8")).hexdigest()[:32]


def bion_trello():
    """Board 'BION – Paid Media'. OneBox pushed its open tasks to the list
    'OneBox' on 30/9; since then the team dragged cards around, Claudia
    created one card by hand, edited a due date and archived another."""
    ex = bion_base()
    board, L = "65f0b1a2c3d4e5f600000001", {
        "onebox": "65f0b1a2c3d4e5f600000011", "curso": "65f0b1a2c3d4e5f600000012",
        "cliente": "65f0b1a2c3d4e5f600000013", "hecho": "65f0b1a2c3d4e5f600000014"}
    ex["project"].update(trelloBoardId=board, trelloBoardName="BION – Paid Media",
                         trelloListId=L["onebox"], trelloListName="OneBox")
    pushed = ["bion-banner", "bion-posts", "bion-urls", "bion-videos", "6c453501da7d40b39c87eb25d3544ae1",
              "9fc4a8e942bb4cc484599e51328adaa0", "88a981daa8e1a54ff0eff6cb7c56f3a7"]
    full = {t["taskId"][:12]: t["taskId"] for t in ex["tasks"]}
    pushed = [full.get(p[:12], p) for p in pushed]
    cards = {}
    for i, tid in enumerate(pushed, 1):
        t = _task(ex, tid)
        cid = f"66a1c0de0000000000000{i:03d}"
        desc = _card_desc(t)
        t.update(trelloCardId=cid, trelloCardListId=L["onebox"], trelloSyncHash=_hash(t["text"], desc, t.get("dueDate", "")),
                 trelloCardUrl=f"https://trello.com/c/{cid[-8:]}")
        cards[tid] = {"id": cid, "name": t["text"], "desc": desc, "due": f"{t['dueDate']}T12:00:00.000Z" if t.get("dueDate") else None,
                      "idList": L["onebox"], "dateLastActivity": "2026-09-30T16:10:00.000Z", "closed": False}
    # What people did in Trello after the push:
    c = cards[full["6c453501da7d"]]                      # due moved by Claudia, now in progress
    c.update(due="2026-10-08T12:00:00.000Z", idList=L["curso"], dateLastActivity="2026-10-02T09:41:12.000Z")
    c = cards["bion-banner"]                              # client says the banner is live
    c.update(idList=L["hecho"], dateLastActivity="2026-10-02T17:05:33.000Z",
             desc=c["desc"] + "\n\nRafael: banner de cookies activo desde el 2/10 (Cookiebot).")
    c = cards["bion-urls"]                                # still waiting, with a note
    c.update(idList=L["cliente"], dateLastActivity="2026-10-01T11:20:00.000Z",
             desc=c["desc"] + "\n\nPadres: lista. Hijos: pendiente de dominio.")
    # Archived to clean the board, NOT because it is done. GET /lists/{id}/cards
    # only returns OPEN cards (Trello's default filter), so it simply vanishes.
    archived = cards.pop(full["88a981daa8e1"])
    c = cards["bion-posts"]                               # moved to done but only 2 of 5 published
    c.update(idList=L["hecho"], dateLastActivity="2026-10-03T19:12:00.000Z",
             desc=c["desc"] + "\n\nRaquel publicó 2 posts (1/10 y 3/10).")
    new_card = {"id": "66a1c0de0000000000000099", "name": "Configurar conversiones de Google Ads (formulario landing Padres)",
                "desc": "Creada en Trello por Claudia. Necesita las URLs definitivas.", "due": "2026-10-10T12:00:00.000Z",
                "idList": L["curso"], "dateLastActivity": "2026-10-02T09:45:00.000Z", "closed": False}
    lists = [{"id": L["onebox"], "name": "OneBox", "closed": False}, {"id": L["curso"], "name": "En curso", "closed": False},
             {"id": L["cliente"], "name": "Esperando cliente", "closed": False}, {"id": L["hecho"], "name": "Hecho", "closed": False},
             {"id": "65f0b1a2c3d4e5f600000019", "name": "Ideas (viejo)", "closed": True}]
    by_list = {lid: [] for lid in L.values()}
    for card in list(cards.values()) + [new_card]:
        by_list[card["idList"]].append(card)
    api = {"_doc": "Raw Trello REST responses, keyed by request. The card of task 88a981da was ARCHIVED: "
                   "the cards endpoint does not return archived cards, so it is absent here.",
           "_archived_card": archived,
           f"GET /1/boards/{board}/lists?fields=name,closed": lists}
    for lid, cs in by_list.items():
        api[f"GET /1/lists/{lid}/cards?fields=name,desc,due,idList,dateLastActivity,closed"] = cs
    _write("bion_trello", "before.json", ex)
    _write("bion_trello", "fixtures/trello_api.json", api)


# ── Farmacia ────────────────────────────────────────────────────────────────
FARMACIA_IDS = {"new-1": "fh-recepcion", "new-2": "fh-entrada-datos", "new-3": "fh-albaranes", "new-4": "fh-estados",
                "new-5": "fh-facturas-mail", "new-6": "fh-facturas-intranet", "new-7": "fh-formulario",
                "new-10": "fh-envios", "new-12": "fh-carpetas-magisfor", "new-13": "fh-acceso-codigo",
                "new-14": "fh-as-is", "new-15": "fh-to-be", "new-16": "fh-pruebas", "new-17": "fh-formacion",
                "new-18": "fh-backlog", "new-19": "fh-pago-30", "new-20": "fh-pago-50", "new-21": "fh-pago-20"}


def farmacia_base():
    # The chatbot, OCR and reverse-engineering tasks were out of the agreed scope: not in the plan.
    ex = _rename(_load("farmacia_reanalyze"), FARMACIA_IDS, drop=("new-8", "new-9", "new-11"))
    ex["project"]["participants"] = [
        {"name": "Jesus", "role": "PM / Dev Backend", "email": "ventas.javc@gmail.com", "phone": ""},
        {"name": "Alejandro Vega", "role": "Consultor automatización", "email": "ale.vega@digitalworkers.example", "phone": ""},
        {"name": "Sandra Prades Hausmann", "role": "Cliente (dueña farmacia)", "email": "sandra@farmacia-hausmann.example", "phone": "+34600000003"},
    ]
    ex["project"]["planVersion"] = 2
    for t in ex["tasks"]:
        if t.get("assignedTo") == "Digital Workers":
            t["assignedTo"] = "Jesus"
    ex["tasks"].append({"projectId": ex["project"]["projectId"], "taskId": "fh-accesos", "status": "blocked",
                        "text": "Sandra debe enviar las contraseñas de la intranet (farmaciahausmann.es)",
                        "assignedTo": "Sandra Prades Hausmann", "blockedReason": "esperando contraseñas",
                        "createdBy": "IA", "dependsOn": []})
    _task(ex, "fh-as-is")["dependsOn"] = ["fh-accesos"]
    _task(ex, "fh-to-be")["dependsOn"] = ["fh-as-is"]
    return ex


ACTA = [
    ("h", "Acta de seguimiento – Proyecto Automatización Farmacia Hausmann (Fase 1)"),
    ("p", "Fecha: 9 de octubre de 2026 · Duración: 45 min · Asistentes: Sandra Prades Hausmann, Jesús Vega, Alejandro Vega"),
    ("h2", "1. Estado"),
    ("p", "El mapeo AS-IS está al 80 %: falta documentar el circuito de devoluciones."),
    ("p", "Sandra entregó las contraseñas de la intranet el 8/10. Ya hay acceso."),
    ("p", "El desarrollador de la intranet confirmó que da acceso al código fuente (correo del 7/10)."),
    ("h2", "2. Decisiones"),
    ("p", "Se prioriza la automatización de facturas (envío por mail y carga en intranet) antes que la recepción de pedidos: "
          "es lo que más horas consume cada decena."),
    ("p", "Objetivo de fin de desarrollo: 14 de noviembre. El pago del 50 % se hace a esa entrega."),
    ("p", "El TO-BE se mantiene para la semana del 13/10."),
    ("h2", "3. Pendientes"),
    ("t", [["Qué", "Quién", "Cuándo"],
           ["Terminar AS-IS (devoluciones)", "Jesús", "10/10"],
           ["Enviar ejemplos de albaranes de envío", "Sandra", "15/10"],
           ["Confirmar si hay datos en las carpetas de Magisfor del servidor", "Sandra", "sin fecha"],
           ["Revisar viabilidad de un informe semanal de pedidos por farmacia", "Alejandro", "si cabe en la fase"]]),
    ("h2", "4. Riesgos"),
    ("p", "El desarrollador de Magisfor deja el soporte en diciembre: si hace falta algo de su lado, tiene que ser antes."),
]


def farmacia_acta():
    from docx import Document
    import io
    ex = farmacia_base()
    _task(ex, "fh-pago-30")["status"] = "done"           # paid before this conversation
    doc = Document()
    for kind, val in ACTA:
        if kind == "h":
            doc.add_heading(val, level=1)
        elif kind == "h2":
            doc.add_heading(val, level=2)
        elif kind == "p":
            doc.add_paragraph(val)
        else:
            tbl = doc.add_table(rows=len(val), cols=len(val[0]))
            for r, row in enumerate(val):
                for c, cell in enumerate(row):
                    tbl.cell(r, c).text = cell
    buf = io.BytesIO()
    doc.save(buf)
    _write("farmacia_acta", "before.json", ex)
    _write("farmacia_acta", "fixtures/acta_seguimiento_2026-10-09.docx", buf.getvalue())


# ── OneBox ──────────────────────────────────────────────────────────────────
ONEBOX_IDS = {"new-1": "ob-verificacion-whatsapp", "new-2": "ob-asana", "new-3": "ob-login", "new-4": "ob-permisos",
              "new-5": "ob-meta-docs", "new-6": "ob-trello-seguimiento", "new-7": "ob-responsive", "new-8": "ob-testing",
              "new-9": "ob-seguridad-chat", "new-10": "ob-fechas-ia", "new-11": "ob-revision-experto",
              "new-12": "ob-graficos", "new-13": "ob-presentacion", "new-14": "ob-microsoft"}


def onebox_base():
    ex = _rename(_load("onebox_chat"), ONEBOX_IDS)
    ex["project"]["participants"].append({"name": "Eloi", "role": "PM externo", "email": "eloi@pm.example", "phone": ""})
    ex["project"]["planVersion"] = 1
    return ex


WEEKLY = """# Update semanal OneBox — semana 1–5 junio

Hola Santi, te resumo la semana:

## Hecho
- **Integración Microsoft / Outlook**: ya conecta, creamos la cuenta Azure el miércoles y probé con mi correo. ✅
- Gráficos del dashboard: listos los de avance y SLA.
- Bug de permisos (usuarios no registrados veían correos generales): corregido y desplegado el jueves.

## En curso
- Seguridad del chatbot: ya filtra por proyecto, falta revisar los teléfonos en las respuestas. Creo que el lunes queda.
- Responsive: arranco la otra semana, primero móvil.

## Bloqueado
- Documentación para Meta (WhatsApp): sigo esperando que Santi me pase el certificado de la empresa.

## Ideas
- Lo de Asana yo lo dejaría para después del MVP, no lo veo prioritario.

Jenny
"""


def onebox_semana():
    ex = onebox_base()
    board, L = "65f0b1a2c3d4e5f600000002", {"backlog": "65f0b1a2c3d4e5f600000021", "sprint": "65f0b1a2c3d4e5f600000022",
                                             "bloq": "65f0b1a2c3d4e5f600000023", "hecho": "65f0b1a2c3d4e5f600000024"}
    ex["project"].update(trelloBoardId=board, trelloBoardName="OneBox MVP", trelloListId="", trelloListName="")
    # Built BY HAND by Jeniffer: no "OneBox task:" ids, names do not match the tasks word for word.
    raw = [
        ("Integración Outlook", "Falta cuenta Azure (Santi).", "2026-06-06", "bloq", "2026-06-02T18:00:00.000Z"),
        ("Seguridad: chatbot no debe mostrar datos de otros proyectos", "Filtrado por proyecto OK. Pendiente teléfonos.",
         "2026-06-08", "sprint", "2026-06-04T20:11:00.000Z"),
        ("Gráficos dashboard (avance, SLA)", "", None, "hecho", "2026-06-03T22:40:00.000Z"),
        ("Fix permisos usuarios no vinculados", "Desplegado.", None, "hecho", "2026-06-04T15:02:00.000Z"),
        ("Responsive móvil / iPad", "", None, "backlog", "2026-05-29T16:00:00.000Z"),
        ("Asana: segunda parte (acciones → tareas)", "", "2026-06-12", "sprint", "2026-05-28T10:00:00.000Z"),
        ("Login con email y contraseña", "Cognito user/pass + MFA", None, "backlog", "2026-05-14T09:00:00.000Z"),
        ("Doc Meta WhatsApp", "Esperando certificado empresa", None, "bloq", "2026-05-21T14:30:00.000Z"),
        ("Exportar informe PDF del proyecto", "Idea de Eloi", None, "backlog", "2026-06-03T11:00:00.000Z"),
    ]
    by_list = {lid: [] for lid in L.values()}
    for i, (name, desc, due, lst, act) in enumerate(raw, 1):
        by_list[L[lst]].append({"id": f"66b2d0de0000000000000{i:03d}", "name": name, "desc": desc,
                                "due": f"{due}T12:00:00.000Z" if due else None, "idList": L[lst],
                                "dateLastActivity": act, "closed": False})
    api = {"_doc": "Raw Trello REST responses, keyed by request. Board made by hand: no OneBox task ids in the cards.",
           f"GET /1/boards/{board}/lists?fields=name,closed": [
               {"id": L["backlog"], "name": "Backlog", "closed": False}, {"id": L["sprint"], "name": "Sprint actual", "closed": False},
               {"id": L["bloq"], "name": "Bloqueado", "closed": False}, {"id": L["hecho"], "name": "Hecho ✅", "closed": False}]}
    for lid, cs in by_list.items():
        api[f"GET /1/lists/{lid}/cards?fields=name,desc,due,idList,dateLastActivity,closed"] = cs
    _write("onebox_semana", "before.json", ex)
    _write("onebox_semana", "fixtures/trello_api.json", api)
    _write("onebox_semana", "fixtures/update_semanal_jenny_2026-06-05.md", WEEKLY)


def plain_bases():
    _write("bion_chat", "before.json", bion_base())
    ex = farmacia_base()
    _write("farmacia_chat", "before.json", ex)
    _write("onebox_lider", "before.json", onebox_base())


if __name__ == "__main__":
    plain_bases()
    bion_trello()
    farmacia_acta()
    onebox_semana()
    print("fixtures written under", HERE)
